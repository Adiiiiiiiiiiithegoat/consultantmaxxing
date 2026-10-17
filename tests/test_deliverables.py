"""Phase 6 gate (offline): number traceability, live spreadsheet formulas, action titles."""

import json
import re

import pytest
from openpyxl import load_workbook

from consultant.agents.storyliner import Argument, Draft, Slide
from consultant.deliverables import model_xlsx, render
from consultant.deliverables.validate import unsupported_numbers, validate_draft
from consultant.research.fact_store import Assumption, Fact
from consultant.workspace import Workspace


def workspace(tmp_path):
    ws = Workspace.create("d1", root=tmp_path, problem_statement="Profits fell 20%; why?",
                          scope={"restated_question": "Why did profits fall 20%?"})
    base = dict(engagement_id="d1", source_type="client", credibility=5,
                source_url="case://x/interviewer")
    ws.facts.add(Fact(claim="Base costs", value=19, unit="crore", quote="costs were Rs 19 crore",
                      **base))                                              # F1
    ws.facts.add(Fact(claim="Current costs", value=22, unit="crore",
                      quote="costs rose to Rs 22 crore", **base))           # F2
    ws.facts.add_assumption(Assumption(engagement_id="d1", name="price_rise", value=0.1,
                                       rationale="match rivals' 12% partly"))  # S1
    ws.run_analysis("cost increase", "size the cost rise",
                    "RESULT = {'increase_cr': 22 - 19, 'increase_pct': (22 - 19) / 19 * 100, "
                    "'irr': 0.1306}",
                    formulas=[{"name": "increase_cr", "expression": "F2 - F1", "unit": "crore"},
                              {"name": "increase_pct", "expression": "increase_cr / F1 * 100",
                               "unit": "%"}])                               # A1
    return ws


def good_draft():
    return Draft(answer="Costs rose 15.79% while revenue was flat; raise prices about 10%.",
                 arguments=[Argument(title="Costs rose by Rs 3 crore, from 19 to 22",
                                     text="That is 15.79%.", refs=["F1", "F2", "A1"]),
                            Argument(title="A 10% price rise closes most of the gap",
                                     text="Estimate.", refs=["S1"])],
                 risks=["Price-sensitive buyers may leave"], next_steps=["Test prices"],
                 open_questions=[], confidence="medium",
                 slides=[Slide(action_title="Input costs, not demand, explain the 20% fall",
                               points=["costs from 19 to 22 crore"], refs=["F1", "F2"])])


def test_formulas_recalculate_to_sandbox_values_and_are_live(tmp_path):
    ws = workspace(tmp_path)
    assert ws.state.analyses[0].formula_issues == []
    path = tmp_path / "model.xlsx"
    stats = model_xlsx.build(ws, path)
    assert stats == {"formulas": 2, "inputs": 3}
    assert model_xlsx.check(path) == []
    model = load_workbook(path)["Model"]
    assert str(model["C2"].value).startswith("=Inputs!")             # a real formula
    live = model_xlsx.recalculate(path)
    assert abs(live[("Model", "C3")] - 300 / 19) < 1e-9
    wb = load_workbook(path)                                           # change an input...
    wb["Inputs"]["C3"] = 25                                            # F2: 22 -> 25
    wb.save(path)
    assert abs(model_xlsx.recalculate(path)[("Model", "C2")] - 6) < 1e-9  # ...model follows
    assert any("increase_cr" in i for i in model_xlsx.check(path))     # and differs from sandbox


def test_formula_mismatch_is_caught_at_analysis_time(tmp_path):
    ws = workspace(tmp_path)
    rec = ws.run_analysis("wrong", "x", "RESULT = {'gap': 4}",
                          formulas=[{"name": "gap", "expression": "F2 - F1"}])
    assert rec.formula_issues and "sandbox gave 4" in rec.formula_issues[0]


def test_validator_traces_numbers_to_cited_evidence(tmp_path):
    ws = workspace(tmp_path)
    assert validate_draft(ws, good_draft()) == []
    bad = good_draft()
    bad.arguments[0].text = "Market share fell 12% and costs rose 15.79%."
    bad.slides.append(Slide(action_title="Overview", points=[], refs=[]))
    bad.arguments[1].refs = ["S1", "F99"]
    kinds = {(i["kind"], i["where"]) for i in validate_draft(ws, bad)}
    assert ("unsupported_number", "argument 1") in kinds        # 12 isn't in F1/F2/A1
    assert ("no_action_title", "slide 2") in kinds
    assert any(k == "unknown_ref" for k, _ in kinds)


def test_number_matching_rules():
    pool = [0.1579, 1536700.0, 22.0]
    assert unsupported_numbers("rose 15.79% (or 15.8%)", pool) == []     # fraction <-> percent
    assert unsupported_numbers("in 2024 and FY2023 it hit 22", pool) == []  # years ignored
    assert unsupported_numbers("1. first point 2. second", [9.0]) == []     # list numbers
    assert unsupported_numbers("about 23", pool) == ["23"]


def test_render_writes_every_deliverable(tmp_path):
    ws = workspace(tmp_path)
    out = render(ws, good_draft())
    for key in ("memo", "storyline", "storyline_md", "model", "sources", "assumptions",
                "validation"):
        assert (ws.dir / out[key]).exists(), key
    report = json.loads((ws.dir / out["validation"]).read_text(encoding="utf-8"))
    assert report["passed"], report
    memo = (ws.dir / out["memo"]).read_text(encoding="utf-8")
    assert memo.index("## Answer") < memo.index("## Why") and "Validation flags" not in memo
    sources = (ws.dir / out["sources"]).read_text(encoding="utf-8")
    assert "| F1 |" in sources and "5/5" in sources
    assert "ASSUMPTION" not in sources and "price_rise" in \
        (ws.dir / out["assumptions"]).read_text(encoding="utf-8")


def test_render_flags_untraceable_numbers_in_the_memo(tmp_path):
    ws = workspace(tmp_path)
    d = good_draft()
    d.answer = "Profits fell because market share dropped 12%."
    render(ws, d)
    memo = (ws.dir / "deliverables/memo.md").read_text(encoding="utf-8")
    assert "Validation flags" in memo and "12" in memo.split("Validation flags")[1]
    assert not json.loads((ws.dir / "deliverables/validation.json").read_text())["passed"]


def test_engagement_validator_forces_revision_when_partner_approves(tmp_path):
    from tests.test_engagement import ScriptedLLM, engagement

    class Untraceable(ScriptedLLM):
        drafts = 0

        def Draft(self, user):
            Untraceable.drafts += 1
            d = super().Draft(user)
            if Untraceable.drafts == 1:
                d["answer"] = "Costs rose 47% (a number nobody gave us)."
            return d

        def Review(self, user):                 # partner always approves
            self.reviews += 1
            return {"issues": [], "strongest_counterargument": "", "ceo_questions": [],
                    "agrees_with_client_prior": False, "prior_agreement_justified": True,
                    "verdict": "approve"}
    llm = Untraceable()
    eng, ws, _ = engagement(tmp_path, llm, eid="val")
    out = eng.run()
    assert Untraceable.drafts == 2 and "story1" in ws.state.steps_done
    assert "47" not in out["draft"].answer
    r0 = ws.state.notes["review0"]
    assert r0["verdict"] == "revise" and "[validator]" in r0["issues"][0]["detail"]


def test_enumeration_only_skips_counting_sequences():
    assert unsupported_numbers("Revenue grew by 2. Then it fell.", [9.0]) == ["2"]
    # smoke-asianpaints-2 false positives: "(1) ... (2) ..." lists and fiscal-year labels
    assert unsupported_numbers("prioritize (1) pricing, (2) mix, (3) costs", [9.0]) == []
    assert unsupported_numbers("FY22‑23 vs FY24-25, FY'24 and FY2023-24", [9.0]) == []
    assert unsupported_numbers("FY24 margin was 23", [9.0]) == ["23"]  # data still checked


def test_formula_that_hardcodes_an_input_is_rejected(tmp_path):
    ws = workspace(tmp_path)
    rec = ws.run_analysis("typed", "x", "RESULT = {'gap': 3, 'pct': 15.79}",
                          formulas=[{"name": "gap", "expression": "22.0 - 19.0"},
                                    {"name": "pct", "expression": "gap / F1 * 100"}])
    assert any("hard-codes 22.0; reference F2" in i for i in rec.formula_issues)
    assert not any(i.startswith("pct: hard-codes") for i in rec.formula_issues)  # 100 is fine


def test_validation_report_fails_on_non_live_formulas(tmp_path):
    ws = workspace(tmp_path)
    ws.run_analysis("typed", "x", "RESULT = {'gap': 3}",
                    formulas=[{"name": "gap", "expression": "22 - 19"}])
    render(ws, good_draft())
    report = json.loads((ws.dir / "deliverables/validation.json").read_text())
    assert not report["passed"] and report["formula_issues"]


def test_repair_pass_fixes_or_removes_every_flagged_number(tmp_path):
    # smoke-asianpaints-2 shipped "43.6" (in no evidence) and 42.39 cited to the wrong facts
    from consultant.deliverables.repair import UNVERIFIED, repair
    from consultant.llm import StructuredOutputError
    from consultant.tracing import BudgetExceeded

    class FakeLLM:
        def __init__(self, reply):
            self.reply, self.calls = reply, 0

        def structured(self, role, msgs, schema):
            self.calls += 1
            if isinstance(self.reply, Exception):
                raise self.reply
            return schema(**self.reply)
    ws = workspace(tmp_path)
    bad = good_draft()
    bad.arguments[1].title = "A 10% price rise closes the gap, about 43.6% of it"  # made up
    bad.slides[0].points = ["costs from 19 to 22 crore (15.79%)"]                 # A1 not cited
    assert {i["where"] for i in validate_draft(ws, bad)} == {"argument 2", "slide 1"}

    # 1) the model re-cites the slide correctly, tries to edit an unflagged location and
    #    leaves the made-up number: only flagged parts change, the leftover is removed
    llm = FakeLLM({"fixes": [
        {"where": "slide 1", "refs": ["F1", "F2", "A1"]},
        {"where": "argument 1", "title": "Rewritten!"},
        {"where": "argument 2", "text": "Still 43.6%."}]})
    d, removed = repair(llm, ws, bad, "evidence")
    assert d.slides[0].refs == ["F1", "F2", "A1"] and "15.79%" in d.slides[0].points[0]
    assert d.arguments[0].title == bad.arguments[0].title
    assert "43.6" not in d.arguments[1].title + d.arguments[1].text
    assert UNVERIFIED in d.arguments[1].text and removed == [
        "43.6 in argument 2 (not in the evidence it cites)"]
    assert validate_draft(ws, d) == []

    # 2) no budget (or no usable reply): every flagged number is removed, nothing else
    for err in (BudgetExceeded("tokens"), StructuredOutputError("x")):
        d, removed = repair(FakeLLM(err), ws, bad, "evidence")
        assert validate_draft(ws, d) == [] and len(removed) == 2
        assert d.slides[0].points == [f"costs from 19 to 22 crore ({UNVERIFIED})"]
        assert d.answer == bad.answer

    # 3) a ref that names no evidence ("Open Questions" on the smoke re-run) is dropped
    odd = good_draft()
    odd.arguments[1].refs = ["S1", "Open Questions"]
    d, removed = repair(FakeLLM({"fixes": []}), ws, odd, "e")
    assert d.arguments[1].refs == ["S1"] and validate_draft(ws, d) == [] and removed == []

    # 4) a clean draft makes no call
    llm = FakeLLM({"fixes": []})
    assert repair(llm, ws, good_draft(), "e") == (good_draft(), []) and llm.calls == 0


DEFERRING = [
    # system-train p012, the answer that lost key_insight and recommendation
    "Given the absence of quantitative data on CAPEX, operating costs, timelines, and "
    "government incentives, we cannot conclusively recommend a specific entry mode or "
    "timeline. The client should first collect the missing data before deciding between "
    "partnership, greenfield build, or acquisition.",
    "The client should collect detailed route-level cost data.",
    "It is too early to decide on an entry strategy.",
    "Gather market data and run a feasibility study.",
    "Hold off on a decision until the pilot results are in, then choose a mode."]
COMMITTED = [
    "Raise cake prices 8% and contract butter supply; collect outlet-level price elasticity "
    "data next quarter to refine the rise.",
    "Enter through a joint venture with a mature-node foundry; confidence is low because capex "
    "data is missing, so commission a capex study as the first next step.",
    "Cut fuel costs by refuelling where fuel is cheaper; we cannot yet size the saving "
    "precisely.",
    "Recommendation: exit the bread line. Before the exit, collect supplier contract data to "
    "time it.",
    "Pursue the premium segment; conduct a pricing study in two outlets to confirm the uplift."]


def test_answer_must_commit_and_carry_the_answer_first_fields(tmp_path):
    from pydantic import ValidationError

    from consultant.agents.storyliner import StoryDraft, render_text
    from consultant.deliverables.validate import defers
    assert all(defers(a) for a in DEFERRING) and not any(defers(a) for a in COMMITTED)
    ws = workspace(tmp_path)
    d = good_draft()
    d.answer = DEFERRING[0]
    assert ("deferred_answer", "answer") in {(i["kind"], i["where"]) for i in
                                             validate_draft(ws, d)}
    fields = dict(confidence_reason="client figures agree", key_evidence=["F1", "A1"],
                  would_change=["costs fall back"])
    with pytest.raises(ValidationError):                # the storyliner must fill them in
        StoryDraft(**good_draft().model_dump(exclude=set(fields)))
    with pytest.raises(ValidationError):                # 2-3 pieces of evidence, not 1
        StoryDraft(**{**good_draft().model_dump(), **fields, "key_evidence": ["F1"]})
    full = StoryDraft(**{**good_draft().model_dump(), **fields})
    assert validate_draft(ws, full) == []
    text = render_text(full)
    assert text.index("RECOMMENDATION:") < text.index("CONFIDENCE: medium (client figures")
    assert "RESTS ON: F1, A1" in text and "WOULD CHANGE THE ANSWER: costs fall back" in text
    bad_ref = StoryDraft(**{**full.model_dump(), "key_evidence": ["F1", "F99"]})
    assert any(i["kind"] == "unknown_ref" for i in validate_draft(ws, bad_ref))
    memo = render(ws, full) and (ws.dir / "deliverables/memo.md").read_text(encoding="utf-8")
    assert memo.index("**Recommendation:**") < memo.index("**Confidence: medium.**") \
        < memo.index("## Why") < memo.index("## Sources") < memo.index("rests on: client")
    answer = memo[:memo.index("## Sources")]           # client-facing: no evidence ids
    assert not re.search(r"\b[FAS]\d+\b", answer)


def test_deferring_draft_is_sent_back_and_the_repair_pass_commits_it(tmp_path):
    from consultant.agents.engagement_manager import fallback_draft
    from consultant.deliverables.repair import repair
    from consultant.deliverables.validate import defers
    from tests.test_engagement import ScriptedLLM, engagement

    class Hedger(ScriptedLLM):
        def Draft(self, user):
            d = super().Draft(user)
            if "PARTNER REVIEW" not in user:          # first draft defers
                d["answer"] = "We cannot recommend a pricing move until more data is collected."
            return d
    llm = Hedger()
    eng, ws, _ = engagement(tmp_path, llm, eid="hedge")
    out = eng.run()
    assert llm.calls.count("StoryDraft") == 2 and not defers(out["draft"].answer)
    review0 = ws.state.notes["review0"]
    assert review0["verdict"] == "revise" and any(
        i["kind"] == "non_sequitur" and "defers" in i["detail"] for i in review0["issues"])

    class Fixer:
        def structured(self, role, msgs, schema):
            return schema(fixes=[{"where": "answer", "text": "Raise cake prices toward "
                                  "competitors and contract butter; test prices first."}])
    w2 = workspace(tmp_path / "w2")
    d = good_draft()
    d.answer = DEFERRING[0]
    fixed, removed = repair(Fixer(), w2, d, "evidence")
    assert not defers(fixed.answer) and removed == [] and validate_draft(w2, fixed) == []

    fb = fallback_draft(ws)                            # budget-exhausted fallback commits too
    assert fb.answer.startswith("Best current answer (low confidence)") and not defers(fb.answer)
    assert fb.confidence == "low" and fb.would_change


def test_cited_fact_must_be_for_the_period_the_claim_states(tmp_path):
    # smoke-asianpaints-2: a repair "fixed" an FY24 claim by citing a stockezee figure for the
    # latest quarter (Jul 2026), and the validator couldn't tell
    from consultant.deliverables.repair import UNVERIFIED, repair
    ws = Workspace.create("p1", root=tmp_path, problem_statement="Why did margins change?",
                          scope={"restated_question": "Why did margins change in FY2023-24?"})
    web = dict(engagement_id="p1", source_type="trade", credibility=3,
               source_url="https://www.stockezee.com/stocks/asianpaint/results")

    def fact(value, period, quote, **kw):
        f = Fact(claim=quote, value=value, unit="%", period=period, quote=quote,
                 **{**web, **kw})
        return f"F{ws.facts.add(f)}"
    q2026 = fact(43.59, "Jul 2026", "Gross Margin % | 43.93% | 42.69% | 43.59% |")
    fy24 = fact(43.59, "FY2023-24", "Gross margin for FY2023-24 was 43.59%")
    q4fy24 = fact(43.59, "Q4 FY24", "Gross margin in Q4 FY24 was 43.59%")
    undated = fact(43.59, "", "Gross margin 43.59%")
    client = fact(43.59, "", "our margin is 43.59%", source_url="case://x/interviewer")

    def flags(point, ref):
        d = Draft(answer="Keep pricing ahead of input costs.", risks=["Input costs rise"],
                  arguments=[Argument(title="Pricing protected the margin", text="See slide.",
                                      refs=[ref])], next_steps=["Track input costs"],
                  open_questions=[], confidence="medium",
                  slides=[Slide(action_title="Gross margin rose despite costlier inputs",
                                points=[point], refs=[ref])])
        return [i for i in validate_draft(ws, d) if i["kind"] == "period_mismatch"], d
    bad, d = flags("FY24 gross margin 43.59%", q2026)
    assert bad and bad[0]["numbers"] == ["43.59"] and "Jul 2026" in bad[0]["detail"]
    assert not flags("FY24 gross margin 43.59%", fy24)[0]           # same year: fine
    assert flags("FY24 gross margin 43.59%", q4fy24)[0]             # quarter cited for year
    assert not flags("Q4 FY24 gross margin 43.59%", q4fy24)[0]      # quarter for quarter
    assert not flags("Gross margin was 43.59%", q2026)[0]           # no period claimed
    fy_label = fact(43.59, "Mar 2024", "Gross margin | Mar 2024 | 43.59%")
    assert not flags("FY24 gross margin 43.59%", fy_label)[0]       # 'Mar 2024' = FY24
    assert flags("FY2023-24 gross margin 43.59%", undated)[0]       # web fact, no period
    assert not flags("FY24 gross margin 43.59%", client)[0]         # client said it
    from consultant.tracing import BudgetExceeded

    class NoBudget:
        def structured(self, *a, **k):
            raise BudgetExceeded("tokens")
    fixed, removed = repair(NoBudget(), ws, d, "evidence")          # no budget: remove it
    assert fixed.slides[0].points == [f"FY24 gross margin {UNVERIFIED}"]
    assert removed == ["43.59 in slide 1 (cited evidence is for another period)"]


def test_client_text_has_no_evidence_ids_and_rounds_raw_floats():
    from consultant.deliverables import display_numbers, strip_refs
    assert strip_refs("Outsourcing wins (A4, F5).") == "Outsourcing wins."
    assert strip_refs("Costs rose [F3] while prices froze [F5, A1].") == \
        "Costs rose while prices froze."
    assert strip_refs("In A4, option B leads.") == "In the analysis, option B leads."
    assert strip_refs("A1, A7 and A8 favour option A.") == "the analyses favour option A."
    assert strip_refs("F2 and A5 agree.") == "the evidence agree."
    keep = "FY24 margin; A-grade suppliers; S&P"
    assert strip_refs(keep) == keep
    assert display_numbers("break-even 0.36666666666666664") == "break-even 0.367"
    assert display_numbers("adds 4.295454545454546% or INR 42525.0") == "adds 4.3% or INR 42525"
    assert display_numbers("2.5x in FY2024, v1.0.3") == "2.5x in FY2024, v1.0.3"


def test_validator_accepts_display_rounding_but_not_a_different_number():
    pool = [42525.0, 0.36666666666666664, 4.295454545454546, 1450.0]
    assert unsupported_numbers("about 42,500, or 0.37 (37%), adding 4.3%", pool) == []
    assert unsupported_numbers("about 1,000", pool) == ["1,000"]     # 1 sig. fig.: not 1,450
    assert unsupported_numbers("about 43,600", pool) == ["43,600"]


def test_one_open_questions_list(tmp_path):
    from consultant.deliverables import open_questions
    ws = Workspace.create("oq", root=tmp_path)
    ws.state.open_questions = ["What is the supplier's minimum order quantity?",
                               "What minimum order quantity does the supplier set?",
                               "How long is the typical delay?"]
    d = good_draft().model_copy(update={"open_questions": ["How long is the typical delay?"]})
    assert open_questions(ws, d) == ["How long is the typical delay?",
                                     "What is the supplier's minimum order quantity?"]


def test_storyliner_is_told_to_calibrate_confidence():
    # prompt-level only: whether the model follows it is a live question (DEFERRED_EVALS)
    from consultant.agents.storyliner import STORY_SYS
    assert "low only when a specific missing fact would likely flip" in STORY_SYS
    assert "never in the prose the client reads" in STORY_SYS


def test_list_outputs_of_an_analysis_are_citable():
    """Gate 1: correct numbers held in a list in an analysis RESULT were removed as
    unverified because lists were skipped."""
    from consultant.deliverables.model_xlsx import _flatten
    out = _flatten({"path": [5, 7.5, {"x": 1}], "cum": 2.5, "flag": True, "name": "y"})
    assert out == {"path[0]": 5, "path[1]": 7.5, "path[2].x": 1, "cum": 2.5}
