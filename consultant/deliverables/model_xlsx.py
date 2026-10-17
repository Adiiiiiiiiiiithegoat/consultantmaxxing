"""Model spreadsheet: Inputs (every numeric fact + assumption, with source) and Model (each
analysis output as a live Excel formula referencing Inputs, next to the sandbox value).

`recalculate()` re-evaluates the formulas *from the saved file* (arithmetic subset: cell
refs, numbers, + - * / ^ and parentheses) so tests can prove the workbook reproduces the
sandbox numbers without needing Excel. Outputs with no formula (e.g. IRR) are written as
values and labelled as computed in the sandbox."""

import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from consultant.research.estimate import safe_eval

_TOKEN = re.compile(r"\b(F\d+|S\d+|[A-Za-z_]\w*)\b")
_REF = re.compile(r"(?:'?([A-Za-z]+)'?!)?\$?([A-Z]{1,3})\$?(\d+)")


def _flatten(d, prefix="") -> dict[str, float]:
    out = {}
    for k, v in (d or {}).items():
        key = f"{prefix}{k}"
        if isinstance(v, bool):
            continue
        if isinstance(v, int | float):
            out[key] = float(v)
        elif isinstance(v, dict):
            out |= _flatten(v, key + ".")
        elif isinstance(v, list | tuple):  # e.g. a path of values: each one is citable
            out |= _flatten({f"[{i}]": x for i, x in enumerate(v)}, key)
    return out


def build(ws, path: Path) -> dict:
    eid = ws.state.engagement_id
    wb = Workbook()
    inp = wb.active
    inp.title = "Inputs"
    inp.append(["ID", "Description", "Value", "Unit", "Period", "Source / rationale", "Type"])
    cell_of: dict[str, str] = {}
    for f in ws.facts.facts(eid):
        if f.value is None:
            continue
        inp.append([f"F{f.id}", f.claim[:200], f.value, f.unit, f.period,
                    f.source_url, f"fact ({f.source_type}, credibility {f.credibility})"])
        cell_of[f"F{f.id}"] = f"Inputs!$C${inp.max_row}"
    for a in ws.facts.assumptions(eid):
        inp.append([f"S{a.id}", a.name, a.value, a.unit, "", a.rationale, "ASSUMPTION (estimate)"])
        cell_of[f"S{a.id}"] = f"Inputs!$C${inp.max_row}"

    model = wb.create_sheet("Model")
    model.append(["Analysis", "Output", "Value (live)", "Unit", "Sandbox value", "Note"])
    written = 0
    for rec in ws.state.analyses:
        if not rec.ok:
            continue
        local: dict[str, str] = {}
        sandbox = _flatten(rec.result)
        done = set()
        for row in rec.formulas:
            expr = row["expression"].replace("**", "^")

            def ref(m, local=local, rid=rec.id):
                tok = m.group(1)
                if tok in local:
                    return local[tok]
                if tok in cell_of:
                    return cell_of[tok]
                raise KeyError(f"{rid}: unknown name {tok} in formula")
            try:
                formula = "=" + _TOKEN.sub(ref, expr)
            except KeyError as e:
                model.append([rec.id, row["name"], sandbox.get(row["name"]), row.get("unit", ""),
                              sandbox.get(row["name"]), f"formula not written: {e}"])
                continue
            model.append([rec.id, row["name"], formula, row.get("unit", ""),
                          sandbox.get(row["name"]), f"{rec.name}: {row['expression']}"])
            local[row["name"]] = f"Model!$C${model.max_row}"
            done.add(row["name"])
            written += 1
        for k, v in sandbox.items():
            if k not in done:
                model.append([rec.id, k, v, "", v, f"computed in sandbox ({rec.dir})"])
    for sheet in (inp, model):
        for c in sheet[1]:
            c.font = Font(bold=True)
    readme = wb.create_sheet("README")
    readme.append(["Inputs: every numeric fact (with source URL) and every assumption "
                   "(labelled ESTIMATE). Model: analysis outputs; 'Value (live)' formulas "
                   "reference Inputs, so changing an input recalculates the model. "
                   "'Sandbox value' is what the analysis code computed."])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"formulas": written, "inputs": len(cell_of)}


def recalculate(path: Path) -> dict[tuple[str, str], float]:
    """Evaluate every formula cell from the file itself. Returns {(sheet, cell): value}."""
    wb = load_workbook(path)  # formulas, not cached values
    cache: dict[tuple[str, str], float] = {}

    def value(sheet: str, addr: str, depth=0) -> float:
        key = (sheet, addr)
        if key in cache:
            return cache[key]
        v = wb[sheet][addr].value
        if isinstance(v, str) and v.startswith("="):
            if depth > 50:
                raise ValueError("formula nesting too deep (cycle?)")
            names = {}

            def sub(m):
                ref_sheet = m.group(1) or sheet
                name = f"r{len(names)}"
                names[name] = value(ref_sheet, f"{m.group(2)}{m.group(3)}", depth + 1)
                return name
            expr = _REF.sub(sub, v[1:]).replace("^", "**")
            v = safe_eval(expr, names)
        cache[key] = float(v)
        return cache[key]

    for sheet in wb.sheetnames:
        for row in wb[sheet].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    value(sheet, c.coordinate)
    return {k: v for k, v in cache.items() if isinstance(wb[k[0]][k[1]].value, str)
            and str(wb[k[0]][k[1]].value).startswith("=")}


def check(path: Path, rel: float = 1e-6) -> list[str]:
    """Live formulas must recalculate to the sandbox value in the same row."""
    wb = load_workbook(path)
    model = wb["Model"]
    live = recalculate(path)
    issues = []
    for r in range(2, model.max_row + 1):
        f = model.cell(r, 3).value
        if isinstance(f, str) and f.startswith("="):
            got, want = live[("Model", f"C{r}")], model.cell(r, 5).value
            if want is None or abs(got - want) > rel * max(1.0, abs(want)):
                issues.append(f"Model!C{r} ({model.cell(r, 2).value}): recalculates to {got}, "
                              f"sandbox {want}")
    return issues
